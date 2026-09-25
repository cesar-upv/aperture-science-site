"""Run with playwright and python-pptx installed; requires npm run dev."""
from pathlib import Path
from playwright.sync_api import sync_playwright, expect
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

root = Path(__file__).resolve().parents[1]
out = root / 'docs/evidencias'
out.mkdir(parents=True, exist_ok=True)
shots = [
 ('00-estrategia-portada','estrategia',0,'1. Nuestra estrategia'),
 ('01-estrategia-filosofia','filosofia',0,'1.1 Nuestra filosofía'),
 ('02-panorama','panorama',0,'1.2 Nuestro panorama'),
 ('03-objetivos','objetivos',0,'1.3 Hacia dónde vamos'),
 ('04-estrategias','estrategias',0,'1.4 Nuestras estrategias'),
 ('05-accion','accion',0,'1.5 De la estrategia a la acción'),
 ('06-organizacion-estructura','estructura',0,'2. Nuestra organización · Nuestra estructura'),
 ('07-areas','areas',0,'2.2 Nuestras áreas'),
 ('08-equipo-direccion','equipo',0,'2.3 Nuestro equipo · Dirección, Manufactura e Investigación'),
 ('09-equipo-personas-calidad','equipo',480,'2.3 Nuestro equipo · Personas, Calidad y Seguridad'),
 ('10-cultura','forma-de-trabajar',0,'3. Nuestra forma de trabajar · Cultura y liderazgo'),
 ('11-colaboracion','cultura-2',0,'3.2–3.5 Liderazgo, comunicación, motivación y colaboración'),
 ('12-compromiso-estandares','compromiso',0,'4. Nuestro compromiso · Nuestros estándares'),
 ('13-resultados','resultados',0,'4.2 Medimos nuestros resultados'),
 ('14-mejora','mejora',0,'4.3 Seguimiento y mejora'),
]
expected = [
 ['Nuestra estrategia','Nuestra filosofía','Nuestro panorama','Hacia dónde vamos','Nuestras estrategias','De la estrategia a la acción'],
 ['Nuestra organización','Nuestra estructura','Nuestras áreas','Nuestro equipo'],
 ['Nuestra forma de trabajar','Nuestra cultura de trabajo','Cómo lideramos','Cómo nos comunicamos','Cómo impulsamos a nuestro equipo','Trabajamos en equipo'],
 ['Nuestro compromiso','Nuestros estándares','Medimos nuestros resultados','Seguimiento y mejora'],
]
with sync_playwright() as p:
 browser = p.chromium.launch()
 page = browser.new_page(viewport={'width':1440,'height':1080},device_scale_factor=2,reduced_motion='reduce')
 errors=[]
 page.on('pageerror',lambda e: errors.append(str(e)))
 page.goto('http://127.0.0.1:5173/',wait_until='networkidle')
 page.evaluate('document.fonts.ready')
 actual=page.locator('.company-section').evaluate_all('(ss)=>ss.map(s=>[...s.querySelectorAll("h2,h3")].map(h=>h.textContent))')
 assert actual==expected, actual
 assert not page.locator('body').inner_text().__contains__('Análisis FODA')
 assert page.locator('.team-portrait').count()==5
 assert page.evaluate('[...document.querySelectorAll("a[href^=\\"#\\"]")].every(a=>document.getElementById(a.hash.slice(1)))')
 page.get_by_role('button',name='Activar portales',exact=True).click()
 assert 'Conexión establecida' in page.get_by_role('status').inner_text()
 page.get_by_role('button',name='Reiniciar demostración').click()
 assert 'Cámara en espera' in page.get_by_role('status').inner_text()
 page.get_by_role('button',name='Ampliar imagen: Cámaras de pruebas').click()
 assert page.locator('dialog').is_visible()
 page.keyboard.press('Escape')
 assert not page.locator('dialog').is_visible()
 for width in [320,390,768,1024,1440,1920]:
  page.set_viewport_size({'width':width,'height':1080})
  assert page.evaluate('document.documentElement.scrollWidth<=innerWidth'),width
 page.set_viewport_size({'width':320,'height':800})
 page.get_by_role('button',name='Abrir menú').click()
 assert page.locator('#mobile-menu').is_visible()
 page.keyboard.press('Escape')
 assert page.get_by_role('button',name='Abrir menú').evaluate('(e)=>e===document.activeElement')
 page.set_viewport_size({'width':1440,'height':1080})
 page.evaluate('document.documentElement.style.fontSize="200%"')
 assert page.evaluate('document.documentElement.scrollWidth<=innerWidth')
 page.evaluate('document.documentElement.style.fontSize=""')
 # Chapter navigation must retain the parent section's active header link.
 page.locator('.strategy-chapters a[href="#objetivos"]').click()
 expect(page.locator('.desktop-links a[href="#estrategia"]')).to_have_attribute('aria-current', 'location')
 assert page.evaluate('document.querySelector("#objetivos").getBoundingClientRect().top >= document.querySelector(".site-header").getBoundingClientRect().bottom')
 assert page.locator('.cover-coordinate .science-mark').evaluate('(e)=>getComputedStyle(e).animationName')=='none'
 assert len(set(page.locator('.story-panel').evaluate_all('(es)=>es.map(e=>getComputedStyle(e).backgroundColor)')))==5
 for name,id,offset,title in shots:
  page.evaluate('([id,offset])=>{const e=document.getElementById(id);scrollTo({top:scrollY+e.getBoundingClientRect().top-105+offset,behavior:"instant"})}',[id,offset])
  page.wait_for_function('Array.from(document.images).filter(i=>{const r=i.getBoundingClientRect();return r.top<innerHeight&&r.bottom>0}).every(i=>i.complete&&i.naturalWidth>0)')
  page.screenshot(path=str(out/f'{name}.png'),animations='disabled')
 assert not errors,errors
 browser.close()

prs=Presentation();prs.slide_width=Inches(12);prs.slide_height=Inches(9.75)
for name,id,offset,title in shots:
 slide=prs.slides.add_slide(prs.slide_layouts[6])
 slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string('081923')
 box=slide.shapes.add_textbox(Inches(.3),Inches(.08),Inches(11.4),Inches(.5))
 par=box.text_frame.paragraphs[0];par.text=title;par.font.size=Pt(19);par.font.color.rgb=RGBColor.from_string('88DFFF')
 slide.shapes.add_picture(str(out/f'{name}.png'),0,Inches(.7),width=Inches(12),height=Inches(9))
prs.save(out/'Unidad-1-Aperture-Science.pptx')
print('PASS: estructura, enlaces, retratos, interacciones, 6 anchos, texto 200%, navegación por apartados, movimiento reducido, sin errores JS. 15 capturas y PPTX generados.')
